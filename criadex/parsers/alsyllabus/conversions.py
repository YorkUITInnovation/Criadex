"""
AL Syllabus DOCX parser — ported from CriaParse (reference service, being retired).

Converts York University AL-format course syllabi (DOCX) into semantic text nodes
so Ragflow can index them as natural-language paragraphs rather than raw binary.

@author Kiarash Bashokian
"""

import io
import re
from typing import Any, Dict, List, Optional, Tuple, TypedDict

import mammoth
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from docx.text.paragraph import Paragraph


class AlNode(TypedDict):
    node_number: int
    text: str
    type: str
    metadata: Dict[str, Any]


def find_h_level(docx_file: Document) -> List[str]:
    headings: List[str] = []
    for paragraph in docx_file.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            headings.append(paragraph.text.strip())
    return [item for item in headings if item]


def find_sections_paragraphs(sections: List[str], docx_file: Document) -> List[int]:
    section_paragraphs: List[int] = []
    for i in range(len(sections)):
        for j in range(len(docx_file.paragraphs)):
            if sections[i] == docx_file.paragraphs[j].text.strip():
                section_paragraphs.append(j)
    return section_paragraphs


def include_hyperlink(paragraph: Paragraph) -> Tuple[List[str], List[str]]:
    hyperlink_text, hyperlink_url = [], []
    if len(paragraph.hyperlinks) > 0:
        for hyperlink in paragraph.hyperlinks:
            hyperlink_text.append(hyperlink.text)
            hyperlink_url.append(hyperlink.url)
    return hyperlink_text, hyperlink_url


def convert_doc_to_nodes(
    section_paragraphs: List[int],
    docx_file: Document,
    sections: List[str],
) -> List[str]:
    nodes_text: List[str] = []
    nodes_temp: str = ""

    for i in range(len(section_paragraphs) - 1):
        for j in range(section_paragraphs[i] + 1, section_paragraphs[i + 1]):
            hyperlink_text, hyperlink_url = include_hyperlink(docx_file.paragraphs[j])
            for k in range(len(hyperlink_text)):
                if not len(hyperlink_text) > 0:
                    continue
                temp_text = docx_file.paragraphs[j].text.replace(
                    hyperlink_text[k],
                    "[" + hyperlink_text[k] + "](" + hyperlink_url[k] + ")",
                )
                docx_file.paragraphs[j].text = temp_text
            nodes_temp = nodes_temp + docx_file.paragraphs[j].text.strip() + " "
        nodes_text.append("*" + sections[i] + "*\n" + nodes_temp + "\n")
        nodes_temp = ""

    for i in range(section_paragraphs[len(section_paragraphs) - 1], len(docx_file.paragraphs)):
        if i == section_paragraphs[len(section_paragraphs) - 1]:
            nodes_temp = nodes_temp + "*" + docx_file.paragraphs[i].text.strip() + "*\n"
            continue
        hyperlink_text, hyperlink_url = include_hyperlink(docx_file.paragraphs[i])
        for k in range(len(hyperlink_text)):
            if not len(hyperlink_text) > 0:
                continue
            temp_text = docx_file.paragraphs[i].text.replace(
                hyperlink_text[k],
                "[" + hyperlink_text[k] + "](" + hyperlink_url[k] + ")",
            )
            docx_file.paragraphs[i].text = temp_text
        nodes_temp = nodes_temp + docx_file.paragraphs[i].text.strip()

    nodes_text.append(nodes_temp)

    if "Course Information" in nodes_text[0]:
        nodes_text[0] += "The course rubric and number is " + docx_file.paragraphs[0].text.strip() + ".\n"
        nodes_text[0] += "The course title is " + docx_file.paragraphs[1].text.strip() + "."

    return nodes_text


def read_tables_bs4(html_text: str) -> List[pd.DataFrame]:
    soup = BeautifulSoup(html_text, "html.parser")
    tables = soup.find_all("table")
    data_frames = []
    for table in tables:
        rows = table.find_all("tr")
        table_data = []
        for row in rows:
            cols = row.find_all("td")
            cols_text = [col.get_text() for col in cols]
            cols_links = [
                col.find("a")["href"] if col.find("a") and "href" in col.find("a").attrs else ""
                for col in cols
            ]
            cols_with_links = [
                f"[{text}] ({link})" if link else text
                for text, link in zip(cols_text, cols_links)
            ]
            table_data.append(cols_with_links)
        data_frames.append(pd.DataFrame(table_data))
    return data_frames


def read_tables(html_text: str, sections: List[str]) -> Tuple[List[pd.DataFrame], List[str]]:
    doc_tables_df = read_tables_bs4(html_text)

    longest_section_len = len(max(sections, key=len)) if sections else 0
    pattern_close = r"</h\d>\s*<table>"
    closing_indexes = [m.start() for m in re.finditer(pattern_close, html_text)]

    previous_index = [ci - longest_section_len for ci in closing_indexes]
    chunks = [html_text[previous_index[i]: closing_indexes[i] + 12] for i in range(len(closing_indexes))]

    almost_titles = []
    for item in chunks:
        match = re.search(r"<h\d>.*</h", item)
        if match:
            almost_titles.append(match.group())

    table_titles = []
    for i in range(len(almost_titles)):
        table_titles.append(almost_titles[i][4:-3])

    for i in range(len(table_titles)):
        if "</h" in table_titles[i]:
            temp_index = table_titles[i].find("</h") + 9
            table_titles[i] = table_titles[i][temp_index:]

    return doc_tables_df, table_titles


def render_tables_add_to_nodes_text(
    table_titles: List[str],
    nodes_text: List[str],
    doc_table_df: List[pd.DataFrame],
) -> List[str]:
    for idx, title in enumerate(table_titles):
        temp_df: pd.DataFrame = doc_table_df[idx]

        if title == "Tutorials":
            temp_text = "*Tutorials*\n "
            temp_text += (
                "Who your TA is and what your TA's email is, and what your tutorial time and day, "
                "your tutorial room, and your tutorial Zoom address are depends on which tutorial "
                "your are in. If the tutorial information is not provided, please always provide "
                "a conditional answer that includes all possibilities. Example of a proper answer: "
                "'if you are in Tutorial 1, your TA is...; if you are in Tutorial 2, your TA is...; "
                "if you are in Tutorial 3, your TA is...'. \n "
            )
            for j in range(1, len(temp_df)):
                temp_text += (
                    "If you are in Tutorial " + temp_df.iloc[j, 0] +
                    ", your TA (or teaching assistant or tutor or responsible instructor who teaches the tutorial) is " +
                    temp_df.iloc[j, 1] + ".\n "
                )
                temp_text += (
                    "If you are in Tutorial " + temp_df.iloc[j, 0] +
                    ", your tutorial time is " + temp_df.iloc[j, 2] + ".\n "
                )
                temp_text += (
                    "If you are in Tutorial " + temp_df.iloc[j, 0] +
                    ", your tutorial room is " + temp_df.iloc[j, 3] + ".\n "
                )
                temp_text += (
                    "If you are in Tutorial " + temp_df.iloc[j, 0] +
                    ", your Zoom address (or Zoom link) during online sessions is " +
                    temp_df.iloc[j, 4] + " .\n"
                )
            nodes_text.append(temp_text)

        elif "Faculty Members Information" in title:
            temp_text = "*Faculty Members Information*\n "
            for j in range(1, len(temp_df)):
                temp_text += (
                    temp_df.iloc[j, 0] + " is the course's " + temp_df.iloc[j, 1] +
                    " and has the following email address: " + temp_df.iloc[j, 2] +
                    " and has the following office hours (time you can meet or appointment time): " +
                    temp_df.iloc[j, 3] +
                    " and has the following office address or location (where you can meet with your professor or instructor or teacher or TA): " +
                    temp_df.iloc[j, 4] + ".\n "
                )
            nodes_text.append(temp_text)

        elif title == "Summary of Evaluation":
            temp_text = "*Summary of Evaluation*\n "
            temp_text += (
                "This section answers questions about how much an assignment is worth (how much it counts toward the final grade) "
                "and when the assignments are due or have to be submitted or handed in (submission date). \n"
            )
            for j in range(1, len(temp_df)):
                temp_text += (
                    "The " + temp_df.iloc[j, 0] + " is worth " + temp_df.iloc[j, 1] +
                    " of the final grade. In other words, it counts for " + temp_df.iloc[j, 1] +
                    " of the final grade.\n "
                )
                temp_text += (
                    "The " + temp_df.iloc[j, 0] + " is due on " + temp_df.iloc[j, 2] +
                    ". In other words, the deadline or due date or submission date for " +
                    temp_df.iloc[j, 0] + " is " + temp_df.iloc[j, 2] + ".\n "
                )
            nodes_text.append(temp_text)

        elif title == "Grading Equivalence":
            temp_text = "*Grading Equivalence*\n "
            for j in range(1, len(temp_df)):
                temp_text += (
                    temp_df.iloc[j, 0] + " is the same as a grade point of " + temp_df.iloc[j, 1] +
                    ", which falls in the percent range of " + temp_df.iloc[j, 2] +
                    "%, and is described as '" + temp_df.iloc[j, 3] + "'.\n "
                )
            nodes_text.append(temp_text)

        elif title == "Definitions of Standing":
            temp_text = "*Definitions of Standing*\n "
            for j in range(0, len(temp_df)):
                temp_text += (
                    "A grade considered '" + temp_df.iloc[j, 0] +
                    "' means that you have a " + temp_df.iloc[j, 1] + "\n "
                )
            nodes_text.append(temp_text)

        elif title == "Schedule and Readings":
            temp_text = "*Schedule and Readings*\n "
            for j in range(1, len(temp_df)):
                temp_text += (
                    "The topic on " + temp_df.iloc[j, 2] +
                    " is (or is about) '" + temp_df.iloc[j, 0] + "'. In other words, '" + temp_df.iloc[j, 0] +
                    "' is presented in class on " + temp_df.iloc[j, 2] + ".\n "
                )
                if str(temp_df.iloc[j, 1]) == "nan":
                    temp_text += "There are no readings on " + temp_df.iloc[j, 2] + ".\n "
                else:
                    temp_text += (
                        "The reading(s) for the topic called '" + temp_df.iloc[j, 0] + "' on " +
                        temp_df.iloc[j, 2] + " is (are) the following: " + str(temp_df.iloc[j, 1]) + "\n "
                    )
            nodes_text.append(temp_text)

        elif title == "Important Dates":
            temp_text = "*Important Dates*\n "
            for j in range(1, len(temp_df)):
                if "None" in temp_df.iloc[j, 1]:
                    temp_text += "There is no " + temp_df.iloc[j, 0] + ".\n "
                else:
                    temp_text += temp_df.iloc[j, 0] + " is on " + temp_df.iloc[j, 1] + ".\n "
            nodes_text.append(temp_text)

        else:
            temp_text = "*" + title + "*\n "
            nb_rows = len(temp_df)
            nb_columns = len(temp_df.columns)
            for j in range(1, nb_rows):
                temp_text += (
                    "The following " + temp_df.iloc[0, 0].lower() + ": " +
                    temp_df.iloc[j, 0] + " has "
                )
                for k in range(1, nb_columns - 1):
                    temp_text += (
                        "the following " + temp_df.iloc[0, k].lower() + ": " +
                        str(temp_df.iloc[j, k]) + " and has "
                    )
                    temp_text += (
                        "the following " + temp_df.iloc[0, k + 1].lower() + ": " +
                        str(temp_df.iloc[j, k + 1]).strip() + "."
                    )
            nodes_text.append(temp_text)

    return nodes_text


def clean_up(nodes_text: List[str]) -> List[str]:
    nodes_text[0] = nodes_text[0].replace(
        "Course Director:",
        "The course director (or professor or instructor or teacher) for this course is ",
    )
    nodes_text[0] = nodes_text[0].replace("Email:", "\n Your course director's email is ")
    nodes_text[0] = nodes_text[0].replace("Semester:", "\n The current semester (or term) is ")
    nodes_text[0] = nodes_text[0].replace(
        "Lecture time & day:",
        "\n The lecture (or class) is offered on the following day and time: ",
    )
    nodes_text[0] = nodes_text[0].replace(
        "Lecture room:",
        "\n If you're wondering how to get to your lecture, the lecture (or class) takes place in the following classroom (or location): ",
    )
    nodes_text[0] = nodes_text[0].replace(
        "Zoom (Lecture):",
        (
            "\n Some classes may be offered on Zoom or you may have to attend some classes on Zoom only during unforeseen "
            "situations such as snowstorms or the instructor's illness, in which case the Zoom link (or Zoom address) for the lecture will be "
        ),
    )
    nodes_text[0] = nodes_text[0].replace(
        "eClass:",
        "\n There is an eClass site (the course has been uploaded to eClass) and the eClass link (or address or URL) is ",
    )
    nodes_text[0] = nodes_text[0].replace(
        "Office:",
        "\n What is the course director's (or professor's or instructor's or teacher's) office number (or office address)? Where can I meet him or her? The answer is: ",
    )
    nodes_text[0] = nodes_text[0].replace(
        "Office Hours:",
        "\n The course director's (or professor's or instructor's or teacher's) office hours are ",
    )
    nodes_text[0] = nodes_text[0].replace("\t", "")

    tutorials_index: Optional[int] = None
    faculty_members_index: Optional[int] = None
    for index, text in enumerate(nodes_text):
        if "*Tutorials*" in text:
            tutorials_index = index
        if "*Faculty Members Information*" in text:
            faculty_members_index = index

    if tutorials_index is not None and faculty_members_index is not None:
        nodes_text[tutorials_index] += nodes_text[faculty_members_index]
        del nodes_text[faculty_members_index]

    nodes_text = [
        text for text in nodes_text
        if not (text.endswith("*\n\n") or text.endswith("*\n \n"))
    ]

    sorted_nodes_text = sorted(nodes_text)

    for i in range(len(sorted_nodes_text) - 2, -1, -1):
        first_index = sorted_nodes_text[i].find("*")
        second_index = sorted_nodes_text[i].find("*", first_index + 1)
        temp_node = sorted_nodes_text[i][:second_index]
        if temp_node == sorted_nodes_text[i + 1][:second_index]:
            sorted_nodes_text[i] = sorted_nodes_text[i] + sorted_nodes_text[i + 1][second_index + 1:]
            del sorted_nodes_text[i + 1]

    return sorted_nodes_text


def convert_to_dict(sorted_nodes_text: List[str], include_ext_metadata_note: bool = False) -> List[AlNode]:
    nodes: List[AlNode] = []
    for node_number, text in enumerate(sorted_nodes_text):
        node_metadata: Dict[str, Any] = {}
        if include_ext_metadata_note:
            node_metadata["al_ext_note"] = "Al Parser Extension Node"
        nodes.append(
            {
                "node_number": node_number,
                "type": "NarrativeText",
                "text": text,
                "metadata": node_metadata,
            }
        )
    return nodes


def convert_file(file_bytes: io.BytesIO) -> List[AlNode]:
    """Convert an AL Syllabus DOCX to semantic text nodes."""
    docx_file: Document = Document(file_bytes)
    html_text: str = mammoth.convert_to_html(file_bytes).value
    sections: List[str] = find_h_level(docx_file)
    section_paragraphs = find_sections_paragraphs(sections, docx_file)
    nodes_text = convert_doc_to_nodes(section_paragraphs, docx_file, sections)
    doc_tables_df, table_titles = read_tables(html_text, sections)
    render_tables_add_to_nodes_text(table_titles, nodes_text, doc_tables_df)
    sorted_nodes_text: List[Any] = clean_up(nodes_text)
    return convert_to_dict(sorted_nodes_text)


def _render_course_information(docx_file: Document, sections: List[str]) -> List[str]:
    """Extract just the Course Information section."""
    if not sections or sections[0] != "Course Information":
        return []

    scrape_text: bool = False
    nodes_text: List[str] = ["*Course Information*\n"]

    for paragraph in docx_file.paragraphs:
        if sections[0] in paragraph.text:
            scrape_text = True
            continue
        if len(sections) > 1 and sections[1] in paragraph.text:
            break
        if scrape_text:
            text: str = paragraph.text.strip().replace("\t", " ")
            if not text:
                continue
            nodes_text[0] += text + "\n"
            nodes_text.append(text)

    nodes_text[0] += "The course rubric and number is " + docx_file.paragraphs[0].text.strip() + ".\n"
    nodes_text[0] += "The course title is " + docx_file.paragraphs[1].text.strip() + "."
    nodes_text[0].strip()
    return nodes_text


def convert_file_partial(file_buffer: io.BytesIO) -> List[AlNode]:
    """Extension mode: parse only AL-syllabus-specific sections (tables + course info).

    Used on top of the generic parser to augment generic content with syllabus-specific
    natural-language renderings of tables.
    """
    file_buffer.seek(0)
    docx_file: Document = Document(file_buffer)
    sections: List[str] = find_h_level(docx_file)

    if len(sections) == 0:
        return []

    file_buffer.seek(0)
    html_text: str = mammoth.convert_to_html(file_buffer).value

    nodes_text: List[str] = _render_course_information(docx_file, sections)
    doc_tables_df, table_titles = read_tables(html_text, sections)
    render_tables_add_to_nodes_text(table_titles, nodes_text, doc_tables_df)
    sorted_nodes_text: List[Any] = clean_up(nodes_text)
    return convert_to_dict(sorted_nodes_text, include_ext_metadata_note=True)
